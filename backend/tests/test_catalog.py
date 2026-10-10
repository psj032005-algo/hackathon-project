import base64
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from main import app


TEST_SECRET = "test-secret-for-isolated-catalog-suite-0123456789"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("STOREFRONT_DATABASE_PATH", str(tmp_path / "catalog-test.sqlite3"))
    monkeypatch.setenv("JWT_SECRET", TEST_SECRET)
    with TestClient(app) as test_client:
        yield test_client


def owner(client, email, slug):
    response = client.post(
        "/api/auth/register",
        json={
            "email": email,
            "password": "a-strong-test-password",
            "store_name": f"{slug} store",
            "store_slug": slug,
        },
    )
    assert response.status_code == 201
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def product_payload(**overrides):
    payload = {
        "name": "Handmade Ceramic Bowl",
        "description": "A glazed bowl made in small batches.",
        "category": "Home",
        "price": 32.5,
        "original_price": 38,
        "discount_percent": 10,
        "stock": 8,
        "sku": "BOWL-001",
        "images": ["https://images.example.test/bowl.jpg"],
        "variants": [{"name": "Color", "value": "Sand", "price_adjustment": 2, "stock": 3}],
    }
    payload.update(overrides)
    return payload


def test_product_crud_filters_and_validation_are_store_scoped(client):
    first_headers = owner(client, "one@example.com", "one-shop")
    second_headers = owner(client, "two@example.com", "two-shop")
    url = "/api/owner/stores/one-shop/products"

    assert client.get(url).status_code == 401
    assert client.post(url, json=product_payload()).status_code == 401
    invalid = client.post(url, headers=first_headers, json=product_payload(price=-1, stock=-3))
    assert invalid.status_code == 422
    assert client.post(url, headers=first_headers, json=product_payload(discount_percent=101)).status_code == 422
    assert client.post(url, headers=first_headers, json=product_payload(images=["javascript:alert(1)"])).status_code == 422
    assert client.post(url, headers=first_headers, json=product_payload(variants=[{"name": "Size", "value": "XL", "stock": -1}])).status_code == 422

    created = client.post(url, headers=first_headers, json=product_payload())
    assert created.status_code == 201
    item = created.json()
    assert item["sku"] == "BOWL-001"
    assert item["images"] == ["https://images.example.test/bowl.jpg"]
    assert item["variants"][0]["value"] == "Sand"

    assert client.get(url.replace("one-shop", "two-shop"), headers=second_headers).json() == []
    assert client.get(url, headers=second_headers).status_code == 404
    assert client.get(url + "?search=BOWL-001", headers=first_headers).json()[0]["id"] == item["id"]
    assert client.get(url + "?category=Home&stock=in_stock", headers=first_headers).json()[0]["id"] == item["id"]
    assert client.get(url + "?stock=low_stock", headers=first_headers).json()[0]["id"] == item["id"]
    assert client.get(url + "?stock=out_of_stock", headers=first_headers).json() == []

    assert client.put(
        f"{url}/{item['id']}", headers=first_headers,
        json=product_payload(name="Updated Ceramic Bowl", stock=12, sku="BOWL-001"),
    ).status_code == 200
    assert client.put(
        f"/api/owner/stores/two-shop/products/{item['id']}",
        headers=second_headers,
        json=product_payload(name="Hijacked"),
    ).status_code == 404
    assert client.delete(f"{url}/{item['id']}", headers=second_headers).status_code == 404
    assert client.get(url, headers=first_headers).json()[0]["name"] == "Updated Ceramic Bowl"
    assert client.delete(f"{url}/{item['id']}", headers=first_headers).status_code == 204
    assert client.get(url, headers=first_headers).json() == []


def test_categories_custom_categories_and_bulk_updates_are_tenant_scoped(client):
    first_headers = owner(client, "first@example.com", "first-shop")
    second_headers = owner(client, "second@example.com", "second-shop")
    first_categories = "/api/owner/stores/first-shop/categories"
    second_categories = "/api/owner/stores/second-shop/categories"

    assert client.get(first_categories).status_code == 401
    assert client.post(first_categories, headers=first_headers, json={"name": "Garden"}).status_code == 201
    assert "Garden" in [row["name"] for row in client.get(first_categories, headers=first_headers).json()]
    assert "Garden" not in [row["name"] for row in client.get(second_categories, headers=second_headers).json()]
    assert client.post(first_categories, headers=first_headers, json={"name": "Garden"}).status_code == 409

    product_url = "/api/owner/stores/first-shop/products"
    product = client.post(product_url, headers=first_headers, json=product_payload()).json()
    assert client.post(
        "/api/owner/stores/second-shop/products", headers=second_headers,
        json=product_payload(name="Invalid Category", sku="BOWL-002", category="Garden"),
    ).status_code == 422
    bulk_url = "/api/owner/stores/first-shop/products/bulk"
    mixed = client.patch(
        bulk_url, headers=first_headers,
        json={"product_ids": [product["id"], 999999], "changes": {"stock": 0}},
    )
    assert mixed.status_code == 404
    assert client.get(product_url, headers=first_headers).json()[0]["stock"] == 8
    changed = client.patch(
        bulk_url, headers=first_headers,
        json={"product_ids": [product["id"]], "changes": {"stock": 0, "category": "Garden"}},
    )
    assert changed.status_code == 200
    assert changed.json()["updated_count"] == 1
    updated = client.get(product_url + "?stock=out_of_stock", headers=first_headers).json()
    assert updated[0]["category"] == "Garden"
    assert updated[0]["variants"][0]["stock"] == 0
    assert client.patch(
        "/api/owner/stores/second-shop/products/bulk", headers=second_headers,
        json={"product_ids": [product["id"]], "changes": {"price": 1}},
    ).status_code == 404


def test_sample_product_import_is_category_based_idempotent_and_tenant_scoped(client):
    first_headers = owner(client, "sample-one@example.com", "sample-one")
    second_headers = owner(client, "sample-two@example.com", "sample-two")
    endpoint = "/api/owner/stores/sample-one/products/import-samples"
    assert client.post(endpoint, json={"categories": ["Home"]}).status_code == 401
    assert client.post(endpoint, headers=second_headers, json={"categories": ["Home"]}).status_code == 404
    invalid = client.post(endpoint, headers=first_headers, json={"categories": ["Not This Store"]})
    assert invalid.status_code == 422
    result = client.post(endpoint, headers=first_headers, json={"categories": ["Home", "Wellness"]})
    assert result.status_code == 200
    assert result.json()["imported_count"] == 4
    products = client.get("/api/owner/stores/sample-one/products", headers=first_headers).json()
    assert {product["category"] for product in products} == {"Home", "Wellness"}
    assert all(product["sku"].startswith("SAMPLE-") and product["images"] for product in products)
    repeated = client.post(endpoint, headers=first_headers, json={"categories": ["Home", "Wellness"]})
    assert repeated.status_code == 200
    assert repeated.json()["imported_count"] == 0 and repeated.json()["skipped_count"] == 4
    assert client.get("/api/owner/stores/sample-two/products", headers=second_headers).json() == []


def test_store_setup_saves_selected_categories_previews_samples_and_completes(client):
    first_headers = owner(client, "setup-one@example.com", "setup-one")
    second_headers = owner(client, "setup-two@example.com", "setup-two")
    setup_url = "/api/owner/stores/setup-one/setup"

    # Mirror the setup page's initial authenticated requests before it renders
    # Step 2. Both endpoints must work against a newly registered store.
    stores = client.get("/api/owner/stores", headers=first_headers)
    assert stores.status_code == 200
    assert [store["slug"] for store in stores.json()] == ["setup-one"]
    categories = client.get("/api/owner/stores/setup-one/categories", headers=first_headers)
    assert categories.status_code == 200
    assert {category["name"] for category in categories.json()} >= {"Home", "Clothing"}

    initial = client.get(setup_url, headers=first_headers)
    assert initial.status_code == 200
    assert initial.json() == {"categories": [], "product_setup_method": None, "is_complete": False}
    assert client.get(setup_url, headers=second_headers).status_code == 404

    setup_payload = {
        # Category matching is case-insensitive, so these aliases all resolve
        # to one Home category row and must not be inserted more than once.
        "categories": [" Home ", "home", "Clothing", "HOME"],
        "product_setup_method": "demo",
    }
    saved = client.put(setup_url, headers=first_headers, json=setup_payload)
    assert saved.status_code == 200
    assert saved.json()["categories"] == ["Home", "Clothing"]
    assert client.get(setup_url, headers=first_headers).json() == {
        "categories": ["Clothing", "Home"], "product_setup_method": "demo", "is_complete": False,
    }

    # Retrying the same request replaces the existing rows atomically rather
    # than colliding with the store/category composite primary key.
    repeated = client.put(setup_url, headers=first_headers, json=setup_payload)
    assert repeated.status_code == 200
    assert client.get(setup_url, headers=first_headers).json()["categories"] == ["Clothing", "Home"]

    # Owners can change their selections on a later submission.
    changed = client.put(setup_url, headers=first_headers, json={
        "categories": ["Wellness"], "product_setup_method": "spreadsheet",
    })
    assert changed.status_code == 200
    assert client.get(setup_url, headers=first_headers).json() == {
        "categories": ["Wellness"], "product_setup_method": "spreadsheet", "is_complete": False,
    }
    restored = client.put(setup_url, headers=first_headers, json=setup_payload)
    assert restored.status_code == 200

    invalid = client.put(setup_url, headers=first_headers, json={
        "categories": ["Unknown"], "product_setup_method": "demo",
    })
    assert invalid.status_code == 422
    # Validation failures must not erase resumable setup progress.
    assert client.get(setup_url, headers=first_headers).json() == {
        "categories": ["Clothing", "Home"], "product_setup_method": "demo", "is_complete": False,
    }

    preview = client.post(
        "/api/owner/stores/setup-one/products/import-samples/preview",
        headers=first_headers, json={"categories": ["Home"]},
    )
    assert preview.status_code == 200
    assert preview.json()["total_products"] == 2
    assert preview.json()["importable_count"] == 2
    assert {row["category"] for row in preview.json()["products"]} == {"Home"}
    not_selected = client.post(
        "/api/owner/stores/setup-one/products/import-samples/preview",
        headers=first_headers, json={"categories": ["Wellness"]},
    )
    assert not_selected.status_code == 422

    assert client.post(setup_url + "/complete", headers=first_headers).status_code == 409
    imported = client.post(
        "/api/owner/stores/setup-one/products/import-samples",
        headers=first_headers, json={"categories": ["Home"]},
    )
    assert imported.status_code == 200
    assert client.post(setup_url + "/complete", headers=first_headers).json() == {"is_complete": True}
    assert client.get(setup_url, headers=first_headers).json()["is_complete"] is True
    assert client.patch("/api/owner/stores/setup-one", headers=first_headers, json={"is_published": True}).status_code == 200
    assert client.get("/api/stores/setup-one/categories").json() == ["Clothing", "Home"]
    public_products = client.get("/api/stores/setup-one/products").json()
    assert len(public_products) == 2
    assert {product["category"] for product in public_products} == {"Home"}


def test_csv_preview_confirmation_row_errors_and_duplicates(client):
    headers = owner(client, "csv@example.com", "csv-shop")
    endpoint = "/api/owner/stores/csv-shop/products/import"
    mapping = {"name": "Product Name", "category": "Type", "price": "Price", "stock": "On Hand", "sku": "SKU", "image_url": "Photo"}
    good_csv = (
        "Product Name,Type,Price,On Hand,SKU,Photo\n"
        "CSV Mug,Home,15.50,4,CSV-1,https://images.example.test/mug.jpg\n"
    )
    request = {"csv_text": good_csv, "mapping": mapping}

    assert client.post(endpoint, headers=headers, json=request).status_code == 400
    preview = client.post(endpoint + "/preview", headers=headers, json=request)
    assert preview.status_code == 200
    assert preview.json()["valid_rows"] == 1
    assert preview.json()["can_import"] is True
    assert client.get("/api/owner/stores/csv-shop/products", headers=headers).json() == []

    imported = client.post(endpoint, headers=headers, json={**request, "confirmed": True})
    assert imported.status_code == 200
    assert imported.json()["imported_count"] == 1
    products = client.get("/api/owner/stores/csv-shop/products", headers=headers).json()
    assert len(products) == 1 and products[0]["sku"] == "CSV-1"

    invalid_csv = "Product Name,Type,Price,On Hand\nBroken,Home,-4,not-a-number\n"
    invalid_request = {"csv_text": invalid_csv, "mapping": {"name": "Product Name", "category": "Type", "price": "Price", "stock": "On Hand"}}
    invalid_preview = client.post(endpoint + "/preview", headers=headers, json=invalid_request)
    assert invalid_preview.status_code == 200
    assert invalid_preview.json()["error_rows"] == 1
    assert len(invalid_preview.json()["rows"][0]["errors"]) >= 2
    rejected = client.post(endpoint, headers=headers, json={**invalid_request, "confirmed": True})
    assert rejected.status_code == 422
    assert len(client.get("/api/owner/stores/csv-shop/products", headers=headers).json()) == 1

    duplicate_csv = "Name,Category,Price\nCSV Mug,Home,18\n"
    duplicate = client.post(
        endpoint + "/preview", headers=headers,
        json={"csv_text": duplicate_csv, "mapping": {"name": "Name", "category": "Category", "price": "Price"}},
    )
    assert duplicate.json()["rows"][0]["valid"] is False
    assert any("Duplicate" in issue for issue in duplicate.json()["rows"][0]["errors"])


def test_csv_limits_malformed_files_and_import_tenant_isolation(client):
    first_headers = owner(client, "csv-one@example.com", "csv-one")
    second_headers = owner(client, "csv-two@example.com", "csv-two")
    url = "/api/owner/stores/csv-one/products/import"
    mapping = {"name": "Name", "category": "Category", "price": "Price"}
    malformed = client.post(
        url + "/preview", headers=first_headers,
        json={"csv_text": 'Name,Category,Price\n"unterminated,Home,4', "mapping": mapping},
    )
    assert malformed.status_code == 422
    too_many_rows = "Name,Category,Price\n" + "X,Home,4\n" * 1001
    limited = client.post(
        url + "/preview", headers=first_headers,
        json={"csv_text": too_many_rows, "mapping": mapping},
    )
    assert limited.status_code == 413
    too_many_bytes = "Name,Category,Price\n" + ("Boîte,Home,4\n" * 80000)
    assert client.post(
        url + "/preview", headers=first_headers,
        json={"csv_text": too_many_bytes, "mapping": mapping},
    ).status_code == 413
    assert client.post(
        url + "/preview", headers=second_headers,
        json={"csv_text": "Name,Category,Price\nOther,Home,3", "mapping": mapping},
    ).status_code == 404


def _xlsx_base64(rows):
    output = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(output)
    workbook.close()
    return base64.b64encode(output.getvalue()).decode("ascii")


def test_xlsx_preview_import_validation_formulas_limits_and_isolation(client):
    first_headers = owner(client, "xlsx-one@example.com", "xlsx-one")
    second_headers = owner(client, "xlsx-two@example.com", "xlsx-two")
    base = "/api/owner/stores/xlsx-one/products/import-xlsx"
    payload = _xlsx_base64([
        ["Product Name", "Type", "Price", "On Hand", "SKU", "Photo"],
        ["Spreadsheet Mug", "Home", 15.5, 4, "XLSX-1", "https://images.example.test/mug.jpg"],
    ])
    mapping = {"name": "Product Name", "category": "Type", "price": "Price", "stock": "On Hand", "sku": "SKU", "image_url": "Photo"}
    assert client.post(base + "/headers", headers=first_headers, json={"xlsx_base64": payload, "mapping": {}}).json()["headers"] == ["Product Name", "Type", "Price", "On Hand", "SKU", "Photo"]
    preview = client.post(base + "/preview", headers=first_headers, json={"xlsx_base64": payload, "mapping": mapping})
    assert preview.status_code == 200 and preview.json()["can_import"] is True
    not_confirmed = client.post(base, headers=first_headers, json={"xlsx_base64": payload, "mapping": mapping})
    assert not_confirmed.status_code == 400
    imported = client.post(base, headers=first_headers, json={"xlsx_base64": payload, "mapping": mapping, "confirmed": True})
    assert imported.status_code == 200 and imported.json()["imported_count"] == 1
    assert client.get("/api/owner/stores/xlsx-one/products", headers=first_headers).json()[0]["sku"] == "XLSX-1"

    formula_payload = _xlsx_base64([["Name", "Category", "Price"], ["Formula", "Home", "=10+5"]])
    formula_result = client.post(base + "/preview", headers=first_headers, json={
        "xlsx_base64": formula_payload,
        "mapping": {"name": "Name", "category": "Category", "price": "Price"},
    })
    assert formula_result.status_code == 422 and "Formula cells" in formula_result.json()["detail"]
    assert client.post(base + "/preview", headers=first_headers, json={"xlsx_base64": "not-base64", "mapping": mapping}).status_code == 422
    assert client.post(base + "/preview", headers=second_headers, json={"xlsx_base64": payload, "mapping": mapping}).status_code == 404
    template = client.get(base + "/template", headers=first_headers)
    assert template.status_code == 200
    assert template.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    template_book = load_workbook(io.BytesIO(template.content), read_only=True, data_only=True)
    assert list(template_book.active.values)[0] == ("name", "description", "category", "price", "original_price", "stock", "sku", "image_url", "discount_percent")
    template_book.close()

    too_many = _xlsx_base64([["Name", "Category", "Price"], *[[f"Item {index}", "Home", 5] for index in range(1001)]])
    assert client.post(base + "/preview", headers=first_headers, json={
        "xlsx_base64": too_many,
        "mapping": {"name": "Name", "category": "Category", "price": "Price"},
    }).status_code == 413


def test_customization_is_validated_persistent_public_and_tenant_scoped(client):
    first_headers = owner(client, "theme-one@example.com", "theme-one")
    second_headers = owner(client, "theme-two@example.com", "theme-two")
    first_url = "/api/owner/stores/theme-one/customization"
    second_url = "/api/owner/stores/theme-two/customization"

    before = client.get(first_url, headers=first_headers)
    assert before.status_code == 200 and before.json()["theme_id"] == "market"
    assert "store_id" not in before.json()
    rejected_store_id = client.patch(first_url, headers=first_headers, json={
        "theme_id": "studio", "store_id": 999999,
    })
    assert rejected_store_id.status_code == 422
    assert any(
        issue["type"] == "extra_forbidden" and issue["loc"][-1] == "store_id"
        for issue in rejected_store_id.json()["detail"]
    )
    assert client.get(first_url, headers=first_headers).json()["theme_id"] == "market"
    changed = client.patch(
        first_url, headers=first_headers,
        json={
            "theme_id": "botanical", "primary_color": "#345678", "accent_color": "#765432",
            "background_color": "#fefefe", "font_family": "serif",
            "banner_url": "https://images.example.test/banner.jpg",
            "homepage_sections": ["hero", "categories"], "footer_text": "Made with care",
            "contact_email": "hello@example.test", "contact_phone": "+1 555 123 4567",
        },
    )
    assert changed.status_code == 200
    assert client.get(first_url, headers=first_headers).json()["primary_color"] == "#345678"
    saved = client.get(first_url, headers=first_headers).json()
    assert saved["accent_color"] == "#765432" and saved["background_color"] == "#fefefe"
    assert saved["font_family"] == "serif"
    assert client.get(second_url, headers=second_headers).json()["primary_color"] != "#345678"
    assert client.patch(first_url, headers=second_headers, json={"theme_id": "midnight"}).status_code == 404
    assert client.patch(first_url, headers=first_headers, json={"primary_color": "red"}).status_code == 422
    assert client.patch(first_url, headers=first_headers, json={"banner_url": "javascript:alert(1)"}).status_code == 422
    assert client.patch(first_url, headers=first_headers, json={"arbitrary_html": "<script>"}).status_code == 422

    assert client.get("/api/stores/theme-one").status_code == 404
    assert client.patch(
        "/api/owner/stores/theme-one", headers=first_headers, json={"is_published": True}
    ).status_code == 200
    public = client.get("/api/stores/theme-one")
    assert public.status_code == 200
    assert public.json()["customization"]["theme_id"] == "botanical"
    assert public.json()["customization"]["homepage_sections"] == ["hero", "categories"]
    for theme in ("market", "botanical", "studio", "midnight"):
        response = client.patch(first_url, headers=first_headers, json={"theme_id": theme})
        assert response.status_code == 200 and response.json()["theme_id"] == theme
        assert client.get("/api/stores/theme-one").json()["customization"]["theme_id"] == theme
